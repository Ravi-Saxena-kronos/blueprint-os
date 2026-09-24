import os
from io import BytesIO
from typing import Optional

from docx import Document

OUTPUT_DIR = os.environ.get("OUTPUT_DIR", "/tmp/blueprint-output")
try:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
except OSError:
    pass


def _build_document(job: dict) -> Document:
    b = job["blueprint"]
    doc = Document()
    doc.add_heading("Blueprint", level=0)
    doc.add_paragraph(f"Job ID: {job['id']}")
    doc.add_paragraph(f"Tier:   {job['tier']}")
    doc.add_paragraph(f"Industry: {job.get('industry') or '—'}")

    pf = b["problem_framing"]
    doc.add_heading("1. Problem framing", level=1)
    doc.add_paragraph(pf["goal"])
    doc.add_paragraph("Stakeholders: " + ", ".join(pf.get("stakeholders", [])))
    doc.add_paragraph("Success criteria:")
    for s in pf.get("success_criteria", []):
        doc.add_paragraph(s, style="List Bullet")

    doc.add_heading("2. Assumptions", level=1)
    for a in b["assumptions"]:
        doc.add_paragraph(f"[{a['id']}] {a['statement']}  ({a['status']})", style="List Bullet")

    doc.add_heading("3. Options", level=1)
    for o in b["options"]:
        doc.add_heading(o["name"], level=2)
        doc.add_paragraph("Pros: " + "; ".join(o["pros"]))
        doc.add_paragraph("Cons: " + "; ".join(o["cons"]))
        doc.add_paragraph("Cost band: " + o["cost_band"])

    doc.add_heading("4. Recommendation", level=1)
    rec = b["recommendation"]
    doc.add_paragraph(rec["chosen_path"])
    doc.add_paragraph(rec["rationale"])
    if rec.get("dissenting_view"):
        doc.add_paragraph("Dissenting view: " + rec["dissenting_view"])

    doc.add_heading("5. Action plan", level=1)
    for s in b["action_plan"]:
        doc.add_paragraph(f"{s['step']}. {s['action']}  — {s['owner_role']}")

    doc.add_heading("6. Risks", level=1)
    for r in b["risks"]:
        doc.add_paragraph(f"{r['description']}  (sev={r['severity']}, lik={r['likelihood']})")
        doc.add_paragraph("  Mitigation: " + r["mitigation"])
        doc.add_paragraph("  Residual:   " + r["residual"])

    doc.add_heading("7. Evidence pack", level=1)
    ep = b["evidence_pack"]
    doc.add_paragraph(f"Confidence per claim: {ep.get('confidence_per_claim')}")
    for s in ep.get("sources", []):
        verified = "verified" if s.get("verified") else "UNVERIFIED"
        doc.add_paragraph(f"- {s['title']} ({verified}) {s.get('url','')}", style="List Bullet")
    doc.add_paragraph("Falsifiers:")
    for x in ep.get("falsifiers", []):
        doc.add_paragraph(x, style="List Bullet")

    doc.add_heading("8. Next steps", level=1)
    for x in b["next_steps"]:
        doc.add_paragraph(x, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph(
        "This document is informational planning assistance with cited sources, "
        "not professional advice and not a warranty of results."
    )
    return doc


def docx_bytes(job: dict) -> bytes:
    bio = BytesIO()
    _build_document(job).save(bio)
    return bio.getvalue()


def render_docx(job: dict) -> str:
    path = os.path.join(OUTPUT_DIR, f"{job['id']}.docx")
    _build_document(job).save(path)
    return path


def maybe_upload_blob(job: dict) -> Optional[str]:
    token = os.environ.get("BLOB_READ_WRITE_TOKEN", "").strip()
    if not token:
        return None
    try:
        import httpx

        data = docx_bytes(job)
        r = httpx.post(
            f"https://blob.vercel-storage.com/blueprints/{job['id']}.docx",
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
            content=data,
            timeout=60.0,
        )
        r.raise_for_status()
        return r.json().get("url")
    except Exception:
        return None
